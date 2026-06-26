import os
import shutil
import pandas as pd

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

def format_report_name(prefix, month, year, suffix=None):

    title_date = f"{month} {year}"

    if suffix:
        return f"{prefix} {suffix} One Calls {title_date}.xlsx"

    return f"{prefix} One Calls {title_date}.xlsx"
# =====================================================
# ENTERPRISE FORMATTING ENGINE (SINGLE SOURCE OF TRUTH)
# =====================================================

def apply_enterprise_style(ws, title, headers_row=2):

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    max_col = ws.max_column if ws.max_column > 0 else 1

    # =====================================================
    # TITLE (DO NOT INSERT ROWS)
    # =====================================================
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.font = Font(size=16, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    title_cell.fill = PatternFill(
        start_color="1F4E79",
        end_color="1F4E79",
        fill_type="solid"
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.row_dimensions[1].height = 28

    # =====================================================
    # HEADER ROW (FIXED POSITION = ROW 2)
    # =====================================================
    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    for cell in ws[2]:   # ALWAYS ROW 2
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = header_fill

    ws.row_dimensions[2].height = 20

    # =====================================================
    # DATA ROWS START AT ROW 3
    # =====================================================
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=max_col):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

    # =====================================================
    # COLUMN WIDTHS
    # =====================================================
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)
# =====================================================
# POTX WORKBOOK
# =====================================================

def build_potx_workbook(contractor, data, writer, month, year):

    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

    title_date = f"{month} {year}"

    # =====================================================
    # SPLIT CLIENT DATA
    # =====================================================
    oxychem = data[
        data["County"]
        .fillna("")
        .str.upper()
        .isin(["HARRIS", "CHAMBERS"])
    ].copy()

    indorama = data[
        data["County"]
        .fillna("")
        .str.upper()
        .eq("ORANGE")
    ].copy()

    # =====================================================
    # FIND TICKET TYPE COLUMN
    # =====================================================
    ticket_type_col = None

    for col in data.columns:

        col_name = str(col).strip().upper()

        if col_name in [
            "TYPE",
            "TICKET TYPE",
            "REQUEST TYPE",
            "WORK TYPE"
        ]:
            ticket_type_col = col
            break

    # =====================================================
    # CREATE EMPTY SUMMARY SHEET
    # =====================================================
    pd.DataFrame().to_excel(
        writer,
        sheet_name="Summary",
        index=False
    )

    ws = writer.sheets["Summary"]

    # =====================================================
    # STYLES
    # =====================================================
    blue_fill = PatternFill(
        start_color="1F4E79",
        end_color="1F4E79",
        fill_type="solid"
    )

    gray_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    divider_border = Border(
        bottom=Side(
            style="medium",
            color="808080"
        )
    )

    # =====================================================
    # HELPER
    # =====================================================
    def write_client_section(start_row, title, df):

        # -----------------------------------------
        # Section Header
        # -----------------------------------------
        ws.merge_cells(
            start_row=start_row,
            start_column=1,
            end_row=start_row,
            end_column=4
        )

        cell = ws.cell(start_row, 1)
        cell.value = title
        cell.font = Font(
            size=14,
            bold=True,
            color="FFFFFF"
        )
        cell.fill = blue_fill
        cell.alignment = Alignment(horizontal="left")

        # -----------------------------------------
        # Total Tickets
        # -----------------------------------------
        total_row = start_row + 2

        ws.cell(total_row, 1, "Total Tickets")
        ws.cell(total_row, 1).font = Font(
            bold=True
        )

        total_cell = ws.cell(total_row, 2)
        total_cell.value = len(df)
        total_cell.font = Font(
            size=16,
            bold=True
        )

        # -----------------------------------------
        # Ticket Type Header
        # -----------------------------------------
        table_row = total_row + 3

        ws.cell(table_row, 1, "Ticket Type")
        ws.cell(table_row, 2, "Count")

        for col in [1, 2]:
            header = ws.cell(table_row, col)
            header.font = Font(bold=True)
            header.fill = gray_fill

        current_row = table_row

        # -----------------------------------------
        # Ticket Type Counts
        # -----------------------------------------
        if ticket_type_col:

            type_counts = (
                df[ticket_type_col]
                .fillna("Unknown")
                .astype(str)
                .str.strip()
                .replace("", "Unknown")
                .value_counts()
            )

            for ticket_type, count in type_counts.items():

                current_row += 1

                ws.cell(
                    current_row,
                    1,
                    ticket_type
                )

                ws.cell(
                    current_row,
                    2,
                    count
                )

        else:

            current_row += 1

            ws.cell(
                current_row,
                1,
                "No Ticket Type Column Found"
            )

            ws.cell(
                current_row,
                2,
                0
            )

        # -----------------------------------------
        # Divider
        # -----------------------------------------
        divider_row = current_row + 2

        for col in range(1, 5):
            ws.cell(
                divider_row,
                col
            ).border = divider_border

        return divider_row + 2

    # =====================================================
    # BUILD DASHBOARD
    # =====================================================
    next_row = 4

    next_row = write_client_section(
        next_row,
        "OXYCHEM",
        oxychem
    )

    next_row = write_client_section(
        next_row,
        "INDORAMA",
        indorama
    )

    # =====================================================
    # DETAIL SHEETS
    # =====================================================
    oxychem.to_excel(
        writer,
        sheet_name="OxyChem",
        index=False,
        startrow=1
    )

    indorama.to_excel(
        writer,
        sheet_name="Indorama",
        index=False,
        startrow=1
    )

    # =====================================================
    # FORMAT SHEETS
    # =====================================================
    for sheet in ["Summary", "OxyChem", "Indorama"]:

        ws = writer.sheets[sheet]

        if sheet == "Summary":
            apply_enterprise_style(
                ws,
                f"POTX Ticket Summary for {title_date}"
            )

        elif sheet == "OxyChem":
            apply_enterprise_style(
                ws,
                f"OxyChem Tickets for {title_date}"
            )

        elif sheet == "Indorama":
            apply_enterprise_style(
                ws,
                f"Indorama (TX Side) for {title_date}"
            )
# =====================================================
# COWBOY WORKBOOK (INTERSECTION DASHBOARD)
# =====================================================

def build_generic_workbook(contractor, data, writer, month, year):

    from openpyxl.styles import (
        Font,
        PatternFill,
        Border,
        Side,
        Alignment
    )

    title_date = f"{month} {year}"

    data = data.copy()

    if "Contractor Code" in data.columns:
        data = data.drop(columns=["Contractor Code"])

    # =====================================================
    # FIND TICKET TYPE COLUMN
    # =====================================================
    ticket_type_col = None

    for col in data.columns:

        col_name = str(col).strip().upper()

        if col_name in [
            "TYPE",
            "TICKET TYPE",
            "REQUEST TYPE",
            "WORK TYPE"
        ]:
            ticket_type_col = col
            break

    # =====================================================
    # CREATE EMPTY SUMMARY SHEET
    # =====================================================
    pd.DataFrame().to_excel(
        writer,
        sheet_name="Summary",
        index=False
    )

    ws = writer.sheets["Summary"]

    # =====================================================
    # STYLES (SAME AS POTX)
    # =====================================================
    blue_fill = PatternFill(
        start_color="1F4E79",
        end_color="1F4E79",
        fill_type="solid"
    )

    gray_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    divider_border = Border(
        bottom=Side(
            style="medium",
            color="808080"
        )
    )

    # =====================================================
    # SECTION HEADER
    # =====================================================
    ws.merge_cells(
        start_row=4,
        start_column=1,
        end_row=4,
        end_column=4
    )

    cell = ws.cell(4, 1)
    cell.value = "COWBOY"
    cell.font = Font(
        size=14,
        bold=True,
        color="FFFFFF"
    )
    cell.fill = blue_fill
    cell.alignment = Alignment(horizontal="left")

    # =====================================================
    # TOTAL TICKETS
    # =====================================================
    ws.cell(6, 1, "Total Tickets")
    ws.cell(6, 1).font = Font(bold=True)

    total_cell = ws.cell(6, 2)
    total_cell.value = len(data)
    total_cell.font = Font(
        size=16,
        bold=True
    )

    # =====================================================
    # TICKET TYPE TABLE
    # =====================================================
    ws.cell(9, 1, "Ticket Type")
    ws.cell(9, 2, "Count")

    for col in [1, 2]:
        header = ws.cell(9, col)
        header.font = Font(bold=True)
        header.fill = gray_fill

    current_row = 9

    if ticket_type_col:

        type_counts = (
            data[ticket_type_col]
            .fillna("Unknown")
            .astype(str)
            .str.strip()
            .replace("", "Unknown")
            .value_counts()
        )

        for ticket_type, count in type_counts.items():

            current_row += 1

            ws.cell(
                current_row,
                1,
                ticket_type
            )

            ws.cell(
                current_row,
                2,
                count
            )

    else:

        current_row += 1
        ws.cell(
            current_row,
            1,
            "No Ticket Type Column Found"
        )
        ws.cell(
            current_row,
            2,
            0
        )

    # =====================================================
    # DIVIDER
    # =====================================================
    divider_row = current_row + 2

    for col in range(1, 5):
        ws.cell(
            divider_row,
            col
        ).border = divider_border

    # =====================================================
    # APPLY ENTERPRISE STYLE
    # =====================================================
    apply_enterprise_style(
        ws,
        f"Cowboy Ticket Summary for {title_date}"
    )

    # =====================================================
    # TICKET DATABASE
    # =====================================================
    data.to_excel(
        writer,
        sheet_name="Tickets",
        index=False,
        startrow=1
    )

    ws = writer.sheets["Tickets"]

    apply_enterprise_style(
        ws,
        f"Cowboy Tickets for {title_date}"
    )
# =====================================================
# CLIENT PACK (OXYCHEM ONLY EXTRACTION)
# =====================================================

def build_client_pack(source_file, base_dir, contractor, month, year):

    client_dir = os.path.join(
        base_dir,
        str(year),
        f"{year}-{month}",
        "Client_Pack"
    )
    os.makedirs(client_dir, exist_ok=True)

    filename = format_report_name(
        prefix=contractor,
        month=month,
        year=year,
        suffix="OxyChem"
    )

    client_file = os.path.join(client_dir, filename)
    # --------------------------------
    # READ POTX MASTER
    # --------------------------------
    wb_source = load_workbook(source_file)

    if "OxyChem" not in wb_source.sheetnames:
        raise ValueError(
            "OxyChem sheet not found in source workbook"
        )

    ws_source = wb_source["OxyChem"]

    # --------------------------------
    # CREATE NEW WORKBOOK
    # --------------------------------
    wb = Workbook()

    ws = wb.active
    ws.title = "OxyChem"

    # --------------------------------
    # COPY EVERYTHING
    # --------------------------------
    for row in ws_source.iter_rows():
        for cell in row:
            ws[cell.coordinate].value = cell.value

    # --------------------------------
    # COPY COLUMN WIDTHS
    # --------------------------------
    for col_letter, dim in ws_source.column_dimensions.items():
        ws.column_dimensions[col_letter].width = dim.width

    # --------------------------------
    # COPY ROW HEIGHTS
    # --------------------------------
    for row_num, dim in ws_source.row_dimensions.items():
        ws.row_dimensions[row_num].height = dim.height

    # --------------------------------
    # APPLY ENTERPRISE STYLING AGAIN
    # --------------------------------
    apply_enterprise_style(
        ws,
        f"OxyChem Tickets for {month} {year}"
    )

    wb.save(client_file)

    return client_file
